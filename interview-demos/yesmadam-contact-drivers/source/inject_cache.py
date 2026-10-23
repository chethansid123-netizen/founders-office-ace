"""Copy LibreOffice-computed values into the openpyxl-built workbook as cached <v> values."""
import datetime as dt
import sys
import zipfile
from lxml import etree
from openpyxl import load_workbook

orig, recalc, out = sys.argv[1:4]
NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
RNS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
vals = load_workbook(recalc, data_only=True)
zin = zipfile.ZipFile(orig)
wbx = etree.fromstring(zin.read("xl/workbook.xml"))
rels = etree.fromstring(zin.read("xl/_rels/workbook.xml.rels"))
rmap = {r.get("Id"): r.get("Target") for r in rels}
sheet_files = {}
for s in wbx.iter(f"{{{NS}}}sheet"):
    t = rmap[s.get(f"{{{RNS}}}id")]
    sheet_files["xl/" + t.lstrip("/").replace("xl/", "")] = s.get("name")
EPOCH = dt.datetime(1899, 12, 30)
filled = missing = 0
zout = zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED)
for item in zin.infolist():
    data = zin.read(item.filename)
    if item.filename in sheet_files:
        ws = vals[sheet_files[item.filename]]
        root = etree.fromstring(data)
        for c in root.iter(f"{{{NS}}}c"):
            f = c.find(f"{{{NS}}}f")
            if f is None:
                continue
            v = ws[c.get("r")].value
            old = c.find(f"{{{NS}}}v")
            if old is not None:
                c.remove(old)
            if v is None:  # formula returned an empty string
                c.set("t", "str"); etree.SubElement(c, f"{{{NS}}}v").text = ""
                missing += 1
                continue
            ve = etree.SubElement(c, f"{{{NS}}}v")
            if isinstance(v, bool):
                c.set("t", "b"); ve.text = "1" if v else "0"
            elif isinstance(v, (int, float)):
                c.attrib.pop("t", None); ve.text = repr(float(v)) if isinstance(v, float) else str(v)
            elif isinstance(v, dt.datetime):
                c.attrib.pop("t", None); ve.text = repr((v - EPOCH).total_seconds() / 86400)
            else:
                c.set("t", "str"); ve.text = str(v)
            filled += 1
        data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    zout.writestr(item, data)
zout.close()
print("cached values written:", filled, "formula cells without a value:", missing)
